"""Import the supplied SCO workbook without inventing reverse ratings.
Usage: python3 scripts/import_sco.py /path/to/workbook.xlsx
Requires openpyxl. Only public/teams is replaced.
"""
import json
from pathlib import Path
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FD = {'Reconnaissance': 'reconnaissance', 'Priority Assets': 'priority-assets',
      'Disruption': 'disruption', 'Take and Hold': 'take-and-hold', 'Purge the Foe': 'purge-the-foe'}
KEYS = ['born-evil', 'team-oasis', 'explosive-sparks', 'king-crimson', 'teamhq',
        'chegongmiao', 'double-disruption', 'blackstone-fortress', 'hogs-of-war',
        'time-fengjun', 'blackstone-obelisk', 'northwatch', 'yueyun-fenghua', 'starmasters']
w = openpyxl.load_workbook(sys.argv[1], data_only=True)
assert len(w.worksheets) == len(KEYS)
teams = []
for key, sheet in zip(KEYS, w):
    players = []
    for i, col in enumerate(range(3, 9)):
        header = sheet.cell(1, col).value
        assert isinstance(header, str)
        name, army = header.split('\n', 1)
        army_list = sheet.cell(13 + i * 3, 1).value
        heading = sheet.cell(12 + i * 3, 1).value
        assert isinstance(heading, str)
        assert heading.startswith(f'{i + 1}. {name} — {army} ('), (sheet.title, heading)
        assert isinstance(army_list, str) and len(army_list) > 100
        players.append({'name': name, 'army': army, 'forceDisposition': FD[sheet.cell(2, col).value],
                        'armyList': army_list, 'scores': {}})
    assert len({p['name'] for p in players}) == 6
    teams.append({'key': key, 'name': sheet.cell(1, 1).value, 'players': players})
home = teams[3]
assert home['name'] == '4. King Crimson'
assert len({p['name'] for t in teams for p in t['players']}) == 84
report = []
for team, sheet in zip(teams, w):
    filled = 0
    for row, player in zip(range(3, 9), home['players']):
        assert sheet.cell(row, 1).value == f"{player['name']} - {player['army']}"
        assert FD[sheet.cell(row, 2).value] == player['forceDisposition']
        for col, opponent in zip(range(3, 9), team['players']):
            score = sheet.cell(row, col).value
            if score is None:
                continue
            assert isinstance(score, (int, float)) and not isinstance(score, bool)
            assert 1 <= score <= 5 and score * 2 == int(score * 2)
            assert opponent['name'] not in player['scores']
            player['scores'][opponent['name']] = score
            filled += 1
    report.append({'team': team['name'], 'filled': filled, 'blank': 36 - filled})
# Validate everything before replacing old data.
out = ROOT / 'public/teams'
out.mkdir(exist_ok=True)
for p in out.glob('*-team.json'):
    p.unlink()
for team in teams:
    (out / f"{team['key']}-team.json").write_text(json.dumps(team, ensure_ascii=False, indent=2) + '\n')
ordered = [home] + [t for t in teams if t is not home]
(out / 'manifest.json').write_text(json.dumps([{'key': t['key'], 'name': t['name']} for t in ordered], ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'teams': len(teams), 'players': sum(len(t['players']) for t in teams),
                  'opponent_scores': sum(x['filled'] for i, x in enumerate(report) if i != 3),
                  'home_practice_scores': report[3]['filled'], 'detail': report}, ensure_ascii=False, indent=2))
